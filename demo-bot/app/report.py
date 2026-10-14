"""Excel-отчёт по заявкам: то, что владелец получает от бота."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .niches import NICHES

STATUS = {"new": "Новая", "ok": "Подтверждена", "no": "Отклонена", "cancel": "Отменена клиентом"}
HEAD = ["№", "Создана", "Компания", "Услуга", "Цена, ₴", "Дата", "Время", "Клиент", "Телефон",
        "Комментарий", "Источник", "Статус", "Запись в CRM"]


def leads_to_xlsx(leads: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
    for l in leads:
        company = NICHES[l["niche"]].company if l["niche"] in NICHES else l["niche"]
        ws.append([l["id"], l["created"].replace("T", " "), company, l["service"], l["price"],
                   l["date"], l["time"], l["name"], l["phone"], l["comment"] or "",
                   l["source"], STATUS.get(l["status"], l["status"]),
                   f"{l['crm']} №{l['external_id']}" if l.get("external_id") else ""])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A2"
    if leads:
        ws.append([])
        ws.append(["", "", "", "Сумма, ₴:", f"=SUM(E2:E{len(leads) + 1})"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
