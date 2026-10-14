"""Логика без Telegram и без ключей. Запуск: python -m tests.test_demo"""
import asyncio
import io
import os
import tempfile
from datetime import date

os.environ["GA4_API_SECRET"] = ""  # тесты не шлют события в настоящий GA4
os.environ["DATA_DIR"] = tempfile.mkdtemp()

from openpyxl import load_workbook  # noqa: E402

from app import ai, cabinet, config, db, i18n  # noqa: E402
from app.bot import lead_card, parse_review, welcome  # noqa: E402
from app.niches import NICHES, find_service, free_slots  # noqa: E402
from app.report import leads_to_xlsx  # noqa: E402
from tests.fakes import ScriptedClient, say, tool  # noqa: E402


def test_data_and_prompt():
    assert len(NICHES) == 4
    for k, n in NICHES.items():
        assert n.services and n.faq and n.slots and n.photo_example, k
        assert any(free_slots(k, d) for d in ai.upcoming_days()), k
        s = ai.static_system(k)
        for phrase in ("Не смог разобрать", "Здравствуйте", "ask_owner", "show_cabinet", "leave_review",
                       "Сейчас уточню у владельца", "Прайс", n.company):
            assert phrase in s, (k, phrase)
        assert "Хочет такого бота" in s            # демо-правило
    assert find_service("salon", "женская стрижка")[0] == 0 and find_service("salon", "ракета") is None
    d = asyncio.run(ai.dynamic_system("salon", {"name": "Анна", "phone": "050"}, []))
    assert "Анна" in d and "не спрашивай" in d and "Свободное время" in d


def test_welcome():
    for k, n in NICHES.items():
        w = welcome(k)
        assert len(w) < 4096 and "Какие услуги есть" in w and "Кнопок нет" in w and "/cabinet" in w
        assert all(s[0] in w for s in n.services) and n.samples[0] in w
        assert "Демо-режим" in w
    config.DEMO_MODE = False
    assert "Демо-режим" not in welcome("salon") and "Хочет такого" not in ai.static_system("salon")
    config.DEMO_MODE = True
    assert all(len(i18n.t(lg, 'how')) < 4096 for lg in i18n.LANGS)


def test_calendar_and_cabinet():
    txt = cabinet.calendar_text(2026, 10, {15}, date(2026, 10, 13))
    lines = txt.split("\n")
    assert "Октябрь 2026" in lines[0] and "[15]" in txt and "(13)" in txt
    assert all(len(l) <= 28 for l in lines) and len(lines) == 2 + 5
    today = date(2026, 10, 14)
    leads = [
        dict(id=1, service="Женская стрижка", price=700, date="2026-10-15", time="10:00", status="ok"),
        dict(id=2, service="Маникюр", price=650, date="2026-11-03", time="12:00", status="new"),
        dict(id=3, service="Укладка", price=500, date="2026-10-01", time="11:00", status="ok"),
        dict(id=4, service="Мелирование", price=2600, date="2026-10-20", time="11:00", status="cancel"),
    ]
    out = cabinet.render({"name": "Анна", "phone": "+380501234567"}, "salon", leads, [], today)
    assert "Анна" in out and "+380501234567" in out and "[15]" in out and "Ноябрь 2026" in out
    assert "№1" in out and "ждёт подтверждения" in out and "подтверждена" in out
    assert "Мелирование" not in out                    # отменённые не показываются
    assert "Визитов: 1" in out and "500 ₴" in out and "до скидки" in out.lower() and "/review" in out
    assert len(out) < 4096
    empty = cabinet.render(None, "auto", [], [], today)
    assert "Профиль пока пустой" in empty and "Ближайших записей нет" in empty
    assert "Анна" in cabinet.profile_text({"name": "Анна"}) and "/name" in cabinet.profile_text(None)


def test_parse_review():
    assert parse_review("5 всё понравилось") == (5, "всё понравилось")
    assert parse_review("4/5 хорошо, но долго") == (4, "хорошо, но долго")
    assert parse_review("Оценка 3: ждал час")[0] == 3
    assert parse_review("⭐⭐⭐ норм")[0] == 3
    assert parse_review("просто отлично")[0] is None
    assert parse_review("10 из 10")[0] is None
    assert parse_review("5 з 5 все чудово") == (5, "все чудово")
    assert parse_review("4 of 5 good but slow") == (4, "good but slow")
    assert parse_review("5 loved it") == (5, "loved it")


async def test_db_and_report():
    await db.init_db()
    await db.touch_user(1, "u")
    await db.set_profile(1, name="Анна")
    await db.set_profile(1, phone="0501234567")
    await db.set_profile(1, name="Ольга")
    u = await db.get_user(1)
    assert u["name"] == "Ольга" and u["phone"] == "0501234567"
    await db.set_niche(1, "auto")
    assert (await db.get_user(1))["niche"] == "auto"
    lid = await db.add_lead(1, "u", "salon", "Женская стрижка", 700, "2026-10-11", "10:00", "Анна", "0501234567",
                            external_id="77", crm="Altegio")
    l = await db.get_lead(lid)
    assert "Анна" in lead_card(l) and "Altegio №77" in lead_card(l)
    await db.set_status(lid, "cancel")
    assert "Отменена" in lead_card(await db.get_lead(lid))
    ws = load_workbook(io.BytesIO(leads_to_xlsx(await db.leads_for(1)))).active
    assert ws["H2"].value == "Анна" and ws["L2"].value == "Отменена клиентом" and ws["M2"].value == "Altegio №77"
    await db.add_review(1, "u", "salon", 5, "супер")
    await db.add_review(1, "u", "salon", 3, "норм")
    st = await db.review_stats()
    assert st["count"] == 2 and st["avg"] == 4.0
    assert len(await db.user_reviews(1, "salon")) == 2 and (await db.stats())["reviews"]["count"] == 2
    assert await db.hit(7, 2, 100) == "ok" and await db.hit(7, 2, 100) == "ok"
    assert await db.hit(7, 2, 100) == "user_limit" and await db.hit(8, 2, 2) == "global_limit"


class FakeActions:
    def __init__(self):
        self.lead_id, self.show_cabinet, self.log = None, False, []

    async def create_lead(self, data):
        self.log.append(("lead", data))
        return 42

    async def update_profile(self, name, phone):
        self.log.append(("profile", name, phone))
        return "Данные профиля обновлены."

    async def cancel_booking(self, bid):
        if bid != 5:
            raise ValueError("Записи с таким номером у этого клиента нет.")
        return "Запись №5 отменена."

    async def leave_review(self, rating, text):
        self.log.append(("review", rating, text))
        return "Отзыв сохранён"

    async def set_language(self, lang):
        self.log.append(("lang", lang))

    async def ask_owner(self, reason, details):
        self.log.append(("owner", reason))
        return "Передано владельцу."


async def test_ai_tools():
    d = ai.upcoming_days()[0]
    t = next(x for x in ai.upcoming_days() if free_slots("salon", x))
    slot = free_slots("salon", t)[0]
    prof = {"name": "Анна", "phone": "0501234567"}

    # запись без имени и телефона: берутся из профиля
    a = FakeActions()
    ai.client = ScriptedClient(tool("create_lead", service="Женская стрижка", date=t.isoformat(), time=slot),
                               say("Записала!"))
    text = await ai.ask(5, "salon", "запишите", None, a, prof, [])
    assert text == "Записала!" and a.lead_id == 42
    assert a.log[0][1]["name"] == "Анна" and a.log[0][1]["phone"] == "0501234567" and a.log[0][1]["price"] == 700
    assert any(b["type"] == "text" and "Анна" in b["text"] for b in ai.client.calls[0]["system"])
    # без профиля и без имени: ошибка модели, записи нет
    a = FakeActions()
    ai.client = ScriptedClient(tool("create_lead", service="Женская стрижка", date=t.isoformat(), time=slot),
                               say("Как вас зовут?"))
    await ai.ask(6, "salon", "запишите", None, a, {}, [])
    assert a.lead_id is None
    err = ai.client.calls[1]["messages"][-1]["content"][0]
    assert err["is_error"] and "имя" in err["content"]
    # занятое время
    a = FakeActions()
    ai.client = ScriptedClient(tool("create_lead", service="Женская стрижка", date=t.isoformat(), time="03:33",
                                    name="Иван", phone="0507654321"), say("Другое время?"))
    await ai.ask(7, "salon", "запишите", None, a, {}, [])
    assert a.lead_id is None and "недоступно" in ai.client.calls[1]["messages"][-1]["content"][0]["content"]
    # остальные инструменты
    a = FakeActions()
    ai.client = ScriptedClient(tool("show_cabinet"), say("Вот кабинет"))
    await ai.ask(8, "salon", "мои записи", None, a, {}, [])
    assert a.show_cabinet
    a = FakeActions()
    ai.client = ScriptedClient(tool("cancel_booking", booking_id=99), say("Не нашла такую запись"))
    await ai.ask(9, "salon", "отмени 99", None, a, {}, [])
    assert ai.client.calls[1]["messages"][-1]["content"][0]["is_error"]
    a = FakeActions()
    ai.client = ScriptedClient(tool("leave_review", rating=5, text="супер"), say("Спасибо"))
    await ai.ask(10, "salon", "5 супер", None, a, {}, [])
    assert a.log == [("review", 5, "супер")]
    a = FakeActions()
    ai.client = ScriptedClient(tool("ask_owner", reason="Просит скидку"), say("Сейчас спрошу у владельца"))
    assert await ai.ask(11, "salon", "дайте скидку 50%", None, a, {}, []) == "Сейчас спрошу у владельца"
    assert a.log == [("owner", "Просит скидку")]
    # смена языка и английское название услуги
    a = FakeActions()
    p2 = {"name": "Анна", "phone": "0501234567"}
    ai.client = ScriptedClient(tool("set_language", lang="en"),
                               tool("create_lead", service="Women's haircut", date=t.isoformat(), time=slot),
                               say("Booked!"))
    await ai.ask(13, "salon", "switch to English and book a haircut", None, a, p2, [])
    assert ("lang", "en") in a.log and p2["lang"] == "en"
    assert a.lead_id == 42 and a.log[-1][1]["service"] == "Женская стрижка"
    a = FakeActions()
    ai.client = ScriptedClient(tool("set_language", lang="de"), say("ok"))
    await ai.ask(14, "salon", "auf Deutsch", None, a, {}, [])
    assert a.log == [] and ai.client.calls[1]["messages"][-1]["content"][0]["is_error"]
    ai.client = ScriptedClient(say("Привіт"))
    await ai.ask(15, "salon", "привіт", None, FakeActions(), {"tg_lang": "uk-UA"}, [])
    assert any("украинский" in b["text"] for b in ai.client.calls[0]["system"])
    assert "uk: Жіноча стрижка; en: Women's haircut" in ai.static_system("salon")
    # фото уходит как image-блок
    ai.client = ScriptedClient(say("Это маникюр"))
    await ai.ask(12, "salon", "", b"\xff\xd8jpeg", FakeActions(), {}, [])
    first = ai.client.calls[0]["messages"][-1]["content"]
    assert first[0]["type"] == "image" and "без подписи" in first[1]["text"]
    assert d  # используется выше


def main():
    test_data_and_prompt()
    test_welcome()
    test_calendar_and_cabinet()
    test_parse_review()
    asyncio.run(test_db_and_report())
    asyncio.run(test_ai_tools())
    print("все проверки пройдены")


if __name__ == "__main__":
    main()
