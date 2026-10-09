"""Читалки: файл -> Doc."""
import csv
import io
import re
import statistics
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from app.docmodel import Block, Doc, make_table, markdown_to_blocks
from app.formats import FormatError

MAX_PDF_PAGES = 100
MAX_IMAGES_FROM_DOC = 30


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


# ---------- TXT / CSV ----------


def read_txt(path: Path) -> Doc:
    doc = Doc(title=path.stem)
    doc.blocks = markdown_to_blocks(_decode(path.read_bytes()))
    return doc


_NUMBER = re.compile(r"^-?(0|[1-9]\d{0,14})([.,]\d+)?$")


def _num(s: str):
    s = s.strip()
    if _NUMBER.match(s):
        try:
            return int(s) if re.fullmatch(r"-?\d+", s) else float(s.replace(",", "."))
        except ValueError:
            return s
    return s


def read_csv(path: Path) -> Doc:
    text = _decode(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        return Doc(title=path.stem)
    body = [[_num(c) for c in r] for r in rows[1:]]
    doc = Doc(title=path.stem)
    doc.blocks = [make_table(rows[0], body, name=path.stem[:31])]
    return doc


# ---------- XLSX ----------


def _cell_value(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) and (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    return v


def read_xlsx(path: Path) -> Doc:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise FormatError(f"не удалось открыть Excel-файл ({e}).")
    doc = Doc(title=path.stem)
    try:
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in r):
                    rows.append([_cell_value(c) for c in r])
                if len(rows) > 5001:
                    doc.notes.append(f"Лист «{ws.title}» обрезан до 5000 строк.")
                    break
            if not rows:
                continue
            width = max(len(r) for r in rows)
            while width > 1 and all((r[width - 1] if len(r) >= width else None) is None for r in rows):
                width -= 1
            rows = [(r + [None] * width)[:width] for r in rows]
            doc.blocks.append(make_table(rows[0], rows[1:], name=ws.title))
    finally:
        wb.close()
    return doc


# ---------- DOCX ----------


def read_docx(path: Path) -> Doc:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = Document(str(path))
    except Exception as e:
        raise FormatError(f"не удалось открыть Word-файл ({e}).")
    doc = Doc(title=path.stem)
    images = 0
    ns_blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    ns_embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            style = (p.style.name if p.style is not None else "") or ""
            if text:
                if style == "Title":
                    doc.title = text
                elif style.startswith("Heading"):
                    m = re.search(r"(\d)", style)
                    level = min(int(m.group(1)), 3) if m else 1
                    doc.blocks.append(Block(f"h{max(level, 1)}", text))
                elif "List" in style:
                    doc.blocks.append(Block("bullet", text))
                else:
                    doc.blocks.append(Block("p", text))
            for blip in child.iter(ns_blip):
                rid = blip.get(ns_embed)
                if rid and images < MAX_IMAGES_FROM_DOC:
                    try:
                        doc.blocks.append(Block("image", image=d.part.related_parts[rid].blob))
                        images += 1
                    except KeyError:
                        pass
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if rows:
                doc.blocks.append(make_table(rows[0], rows[1:]))
    return doc


# ---------- PPTX ----------


def _shapes_in_order(shapes):
    items = sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
    for sh in items:
        if sh.shape_type == 6:  # группа
            yield from _shapes_in_order(sh.shapes)
        else:
            yield sh


def read_pptx(path: Path) -> Doc:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise FormatError(f"не удалось открыть презентацию ({e}).")
    doc = Doc(title=path.stem)
    images = 0
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if n == 1 and title:
            doc.title = title
        else:
            doc.blocks.append(Block("h2", title or f"Слайд {n}"))
        for sh in _shapes_in_order(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        doc.blocks.append(Block("bullet", text))
            elif getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                if rows:
                    doc.blocks.append(make_table(rows[0], rows[1:]))
            elif sh.shape_type == 13 and images < MAX_IMAGES_FROM_DOC:   # картинка
                try:
                    doc.blocks.append(Block("image", image=sh.image.blob))
                    images += 1
                except Exception:
                    pass
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.blocks.append(Block("p", f"Заметки докладчика: {notes}"))
    return doc


# ---------- PDF ----------


def _para_blocks(lines: list[dict], body_size: float) -> list[Block]:
    """Строки страницы -> заголовки, пункты, абзацы."""
    if not lines:
        return []
    maxlen = max(len(l["text"]) for l in lines) or 1
    blocks: list[Block] = []
    para: list[str] = []

    def flush():
        if para:
            blocks.append(Block("p", " ".join(para).strip()))
            para.clear()

    for l in lines:
        text = l["text"].strip()
        if not text:
            continue
        size = l["size"]
        if body_size and size >= body_size * 1.25 and len(text) < 120:
            flush()
            blocks.append(Block("h1" if size >= body_size * 1.6 else "h2", text))
            continue
        if re.match(r"^[•●▪◦\-–]\s+", text):
            flush()
            blocks.append(Block("bullet", re.sub(r"^[•●▪◦\-–]\s+", "", text)))
            continue
        para.append(text)
        if len(text) < maxlen * 0.6 or (text[-1] in ".!?:" and len(text) < maxlen * 0.9):
            flush()
    flush()
    return blocks


def read_pdf(path: Path) -> Doc:
    import pdfplumber

    doc = Doc(title=path.stem)
    try:
        pdf = pdfplumber.open(str(path))
    except Exception:
        raise FormatError("PDF не открывается: он повреждён или защищён паролем.")
    with pdf:
        pages = pdf.pages[:MAX_PDF_PAGES]
        if len(pdf.pages) > MAX_PDF_PAGES:
            doc.notes.append(f"Обработаны первые {MAX_PDF_PAGES} страниц из {len(pdf.pages)}.")

        # самый частый размер шрифта считаем «основным текстом»
        sizes: Counter = Counter()
        page_lines: list[tuple] = []
        had_images = False
        for page in pages:
            had_images = had_images or bool(page.images)
            tables = page.find_tables()
            rest = page
            for t in tables:
                rest = rest.outside_bbox(t.bbox)
            lines = []
            for ln in rest.extract_text_lines(return_chars=True) or []:
                chars = ln.get("chars") or []
                size = statistics.median(c.get("size", 0) for c in chars) if chars else 0
                lines.append({"text": ln["text"], "size": size, "top": ln["top"]})
                sizes[round(size, 1)] += len(ln["text"])
            tbl_blocks = []
            for t in tables:
                rows = t.extract()
                rows = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in r] for r in rows if r]
                if len(rows) >= 2 or (rows and len(rows[0]) > 1):
                    tbl_blocks.append((t.bbox[1], make_table(rows[0], rows[1:])))
            page_lines.append((lines, tbl_blocks))
        body_size = sizes.most_common(1)[0][0] if sizes else 0

        for lines, tbl_blocks in page_lines:
            # таблицы вставляем по вертикальному положению среди текста
            before: list[dict] = []
            pending = sorted(tbl_blocks, key=lambda x: x[0])
            for l in lines:
                while pending and pending[0][0] < l["top"]:
                    doc.blocks += _para_blocks(before, body_size)
                    before = []
                    doc.blocks.append(pending.pop(0)[1])
                before.append(l)
            doc.blocks += _para_blocks(before, body_size)
            for _, tb in pending:
                doc.blocks.append(tb)
        if had_images:
            doc.notes.append("Изображения внутри PDF не перенесены (только текст и таблицы).")
    return doc


# ---------- IMAGE ----------


def read_image(path: Path) -> Doc:
    from PIL import Image

    try:
        img = Image.open(path)
        img.load()
    except Exception:
        raise FormatError("не удалось открыть изображение.")
    buf = io.BytesIO()
    if img.mode not in ("RGB", "RGBA", "L"):
        img = img.convert("RGB")
    img.save(buf, format="PNG")
    doc = Doc(title=path.stem)
    doc.blocks = [Block("image", image=buf.getvalue())]
    return doc


READERS = {
    "docx": read_docx, "xlsx": read_xlsx, "csv": read_csv, "pptx": read_pptx,
    "pdf": read_pdf, "image": read_image, "txt": read_txt,
}
