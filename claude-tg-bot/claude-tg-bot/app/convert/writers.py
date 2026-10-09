"""Писалки: Doc -> xlsx, docx, pptx, pdf, png."""
import io
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from PIL import Image  # noqa: E402

from app.docmodel import Block, Design, Doc, Slide, doc_to_slides, make_table  # noqa: E402
from app.formats import FormatError  # noqa: E402

_FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"      # DejaVu есть в matplotlib, кириллица работает
_REGULAR = _FONT_DIR / "DejaVuSans.ttf"
_BOLD = _FONT_DIR / "DejaVuSans-Bold.ttf"

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return _CTRL.sub("", str(v))


def _img_size(data: bytes) -> tuple[int, int]:
    with Image.open(io.BytesIO(data)) as im:
        return im.size


def blocks_of(doc: Doc) -> list[Block]:
    """Блоки документа; если есть только слайды (презентация от Claude), собираем из них."""
    if doc.blocks or not doc.slides:
        return doc.blocks
    out: list[Block] = []
    for s in doc.slides:
        if s.layout == "title":
            out.append(Block("h1", s.title))
            if s.subtitle:
                out.append(Block("p", s.subtitle))
            continue
        if s.title:
            out.append(Block("h2", s.title))
        out += [Block("bullet", b) for b in s.bullets]
        if s.table is not None:
            out.append(s.table)
        if s.image:
            out.append(Block("image", image=s.image))
        if s.notes:
            out.append(Block("p", f"Заметки: {s.notes}"))
    return out


# ============================ графики ============================


def make_chart_image(chart: dict, title: str = "") -> bytes:
    labels = [_s(x) for x in chart.get("labels", [])]
    series = chart.get("series", [])
    kind = chart.get("type", "bar")
    if not series or not labels:
        raise ValueError("В chart нужны labels и series с данными")
    for s in series:
        if len(s.get("values", [])) != len(labels):
            raise ValueError("В chart число values в каждой серии должно совпадать с числом labels")
    fig, ax = plt.subplots(figsize=(10, 6), dpi=150)
    try:
        if kind == "pie":
            ax.pie(series[0]["values"], labels=labels, autopct="%1.1f%%", startangle=90)
            ax.axis("equal")
        elif kind == "line":
            for s in series:
                ax.plot(labels, s["values"], marker="o", label=s.get("name"))
        else:
            n = len(series)
            width = 0.8 / n
            xs = list(range(len(labels)))
            for i, s in enumerate(series):
                ax.bar([x + i * width - 0.4 + width / 2 for x in xs], s["values"], width, label=s.get("name"))
            ax.set_xticks(xs)
            ax.set_xticklabels(labels)
        if kind != "pie":
            ax.set_xlabel(chart.get("x_label", ""))
            ax.set_ylabel(chart.get("y_label", ""))
            ax.grid(axis="y", alpha=0.3)
            if len(labels) > 6:
                plt.setp(ax.get_xticklabels(), rotation=30, ha="right")
            if any(s.get("name") for s in series):
                ax.legend()
        if title:
            ax.set_title(title, fontsize=14, fontweight="bold")
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        return buf.getvalue()
    finally:
        plt.close(fig)


# ============================ XLSX ============================


def write_xlsx(doc: Doc, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tables = doc.tables
    if not tables:
        raise FormatError("Excel принимает только таблицы, а в содержимом таблиц нет.")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for i, t in enumerate(tables, 1):
        name = re.sub(r"[\\/*?:\[\]]", "", t.name or f"Таблица {i}")[:31] or f"Таблица {i}"
        base, k = name, 1
        while name.lower() in used:
            k += 1
            name = f"{base[:28]}_{k}"
        used.add(name.lower())
        ws = wb.create_sheet(name)
        ws.append([_s(h) for h in t.headers])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in t.rows:
            ws.append([v[:32000] if isinstance(v, str) else v for v in row])
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"      # текст, а не формула
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for ci, col in enumerate(ws.columns, 1):
            width = max((len(_s(c.value)) for c in col), default=8)
            ws.column_dimensions[get_column_letter(ci)].width = min(max(width + 2, 10), 60)
    wb.save(path)


# ============================ DOCX ============================


def _shade(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _docx_runs(par, text: str) -> None:
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", _s(text))):
        if chunk:
            par.add_run(chunk).bold = bool(i % 2)


def write_docx(doc: Doc, path: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    blocks = blocks_of(doc)
    if doc.title and not (blocks and blocks[0].kind == "h1" and blocks[0].text == doc.title):
        d.add_heading(_s(doc.title), 0)
    for b in blocks:
        if b.kind in ("h1", "h2", "h3"):
            d.add_heading(_s(b.text).replace("**", ""), int(b.kind[1]))
        elif b.kind == "bullet":
            _docx_runs(d.add_paragraph(style="List Bullet"), b.text)
        elif b.kind == "p":
            _docx_runs(d.add_paragraph(), b.text)
        elif b.kind == "table":
            if b.name:
                d.add_heading(_s(b.name), 2)
            ncols = max(len(b.headers), 1)
            tbl = d.add_table(rows=1, cols=ncols, style="Table Grid")
            for i, h in enumerate(b.headers):
                cell = tbl.rows[0].cells[i]
                cell.text = ""
                run = cell.paragraphs[0].add_run(_s(h))
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(cell, "2F5597")
            for row in b.rows:
                cells = tbl.add_row().cells
                for i, v in enumerate(row[:ncols]):
                    cells[i].text = _s(v)
            d.add_paragraph()
        elif b.kind == "image" and b.image:
            try:
                w, _h = _img_size(b.image)
                d.add_picture(io.BytesIO(b.image), width=Inches(min(6.2, max(w / 96, 1.5))))
            except Exception:
                doc.notes.append("Одно из изображений не удалось вставить.")
    d.save(path)


# ============================ PDF ============================

_FONTS_READY = False


def _register_fonts() -> None:
    global _FONTS_READY
    if _FONTS_READY:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("DejaVu", str(_REGULAR)))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(_BOLD)))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold")
    _FONTS_READY = True


def _esc(text) -> str:
    return _s(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _inline_html(text) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _esc(text))


def write_pdf(doc: Doc, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image as RLImage, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    _register_fonts()
    blocks = blocks_of(doc)
    wide = any(b.kind == "table" and len(b.headers) > 6 for b in blocks)
    page = landscape(A4) if wide else A4
    margin = 16 * mm
    frame_w = page[0] - 2 * margin
    frame_h = page[1] - 2 * margin - 10 * mm

    ss = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=ss["Normal"], fontName="DejaVu", fontSize=10, leading=14)
    styles = {
        "p": base,
        "h1": ParagraphStyle("h1", parent=base, fontName="DejaVu-Bold", fontSize=16, leading=20, spaceBefore=10, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=base, fontName="DejaVu-Bold", fontSize=13, leading=17, spaceBefore=8, spaceAfter=4),
        "h3": ParagraphStyle("h3", parent=base, fontName="DejaVu-Bold", fontSize=11, leading=15, spaceBefore=6, spaceAfter=3),
    }
    title_style = ParagraphStyle("title", parent=styles["h1"], fontSize=20, leading=24, spaceAfter=12)

    story = []
    if doc.title and not (blocks and blocks[0].kind == "h1" and blocks[0].text == doc.title):
        story.append(Paragraph(_esc(doc.title), title_style))
    bullets: list = []

    def flush_bullets():
        if bullets:
            story.append(ListFlowable([ListItem(b, leftIndent=12) for b in bullets],
                                      bulletType="bullet", start="•", leftIndent=14))
            bullets.clear()

    for b in blocks:
        if b.kind == "bullet":
            bullets.append(Paragraph(_inline_html(b.text), base))
            continue
        flush_bullets()
        if b.kind in ("h1", "h2", "h3", "p"):
            story.append(Paragraph(_inline_html(b.text), styles[b.kind]))
            story.append(Spacer(1, 3))
        elif b.kind == "table":
            ncols = max(len(b.headers), 1)
            fs = 8.5 if ncols <= 8 else 7
            cell_style = ParagraphStyle("cell", parent=base, fontSize=fs, leading=fs + 2.5)
            head_style = ParagraphStyle("head", parent=cell_style, fontName="DejaVu-Bold", textColor=colors.white)
            story.append(Spacer(1, 6))
            if b.name:
                story.append(Paragraph(_esc(b.name), styles["h2"]))
            data = [[Paragraph(_esc(h), head_style) for h in b.headers]]
            for row in b.rows:
                data.append([Paragraph(_esc(v), cell_style) for v in row[:ncols]])
            tbl = Table(data, repeatRows=1, colWidths=[frame_w / ncols] * ncols)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 6))
        elif b.kind == "image" and b.image:
            try:
                w, h = _img_size(b.image)
                scale = min(frame_w / w, frame_h / h, 1.0) if w and h else 1.0
                # маленькие картинки не растягиваем, большие вписываем в страницу
                story.append(RLImage(io.BytesIO(b.image), width=w * scale, height=h * scale))
                story.append(Spacer(1, 6))
            except Exception:
                doc.notes.append("Одно из изображений не удалось вставить.")
    flush_bullets()

    SimpleDocTemplate(
        str(path), pagesize=page, title=_s(doc.title) or "Document",
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
    ).build(story or [Paragraph("(пусто)", base)])


# ============================ PPTX ============================

THEMES = {
    "corporate_blue": dict(bg="FFFFFF", title="1F3864", text="262626", accent="2F5597"),
    "light": dict(bg="FFFFFF", title="111827", text="374151", accent="0F766E"),
    "dark": dict(bg="1E1E2E", title="FFFFFF", text="E5E7EB", accent="7AA2F7"),
    "minimal": dict(bg="FAFAFA", title="000000", text="404040", accent="000000"),
    "warm": dict(bg="FFF8F0", title="7A2E0E", text="3B2A20", accent="D9480F"),
    "green": dict(bg="F4FAF4", title="14532D", text="1F2937", accent="2E7D32"),
}
_HEX = re.compile(r"^#?([0-9a-fA-F]{6})$")


def _hex(v: str | None) -> str | None:
    m = _HEX.match((v or "").strip())
    return m.group(1).upper() if m else None


def _lum(h: str) -> float:
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b


def _mix(a: str, b: str, t: float) -> str:
    """Смесь цветов: t=0 -> a, t=1 -> b."""
    return "".join(f"{round(int(a[i:i+2], 16) * (1 - t) + int(b[i:i+2], 16) * t):02X}" for i in (0, 2, 4))


def resolve_design(d: Design) -> dict:
    theme = dict(THEMES.get(d.theme, THEMES["corporate_blue"]))
    bg = _hex(d.background)
    if bg:
        theme["bg"] = bg
        dark = _lum(bg) < 0.5
        if not _hex(d.text_color):
            theme["text"] = "F3F4F6" if dark else "1F2937"
        if (_lum(theme["title"]) < 0.5) == dark:      # заголовок слился бы с фоном
            theme["title"] = "FFFFFF" if dark else "111827"
    if _hex(d.accent):
        theme["accent"] = _hex(d.accent)
    if _hex(d.text_color):
        theme["text"] = _hex(d.text_color)
    theme["font"] = (d.font or "Calibri").strip()[:40] or "Calibri"
    return theme


def _expand_slides(slides: list[Slide], max_bullets: int = 7, max_rows: int = 12) -> list[Slide]:
    out: list[Slide] = []
    for s in slides:
        if s.layout == "bullets" and len(s.bullets) > max_bullets:
            chunks = [s.bullets[i:i + max_bullets] for i in range(0, len(s.bullets), max_bullets)]
            for n, ch in enumerate(chunks):
                out.append(Slide("bullets", s.title if n == 0 else f"{s.title} (продолжение)", bullets=ch,
                                 notes=s.notes if n == 0 else ""))
        elif s.layout == "table" and s.table is not None and len(s.table.rows) > max_rows:
            rows = s.table.rows
            parts = [rows[i:i + max_rows] for i in range(0, len(rows), max_rows)]
            for n, ch in enumerate(parts, 1):
                out.append(Slide("table", f"{s.title} ({n}/{len(parts)})", table=make_table(s.table.headers, ch),
                                 notes=s.notes if n == 1 else ""))
        else:
            out.append(s)
    return out


def write_pptx(doc: Doc, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    th = resolve_design(doc.design)
    font = th["font"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = _expand_slides(doc_to_slides(doc))
    if not slides:
        raise FormatError("в содержимом нет данных для презентации.")

    def rgb(h):
        return RGBColor.from_string(h)

    def runs(par, text, size, color, bold=False):
        for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", _s(text))):
            if not chunk:
                continue
            r = par.add_run()
            r.text = chunk
            r.font.size = Pt(size)
            r.font.bold = bold or bool(i % 2)
            r.font.name = font
            r.font.color.rgb = rgb(color)

    def textbox(slide, l, t, w, h, text, size, color, bold=False, align=None, anchor=None):
        box = slide.shapes.add_textbox(l, t, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        if align:
            p.alignment = align
        runs(p, text, size, color, bold)
        return box

    def rect(slide, l, t, w, h, color):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def background(slide, color):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(color)

    def title_size(text, big=32):
        n = len(text or "")
        return big if n < 45 else big - 6 if n < 75 else big - 10

    def bullet_para(par, color):
        pPr = par._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.38))))
        pPr.set("indent", str(-Emu(Inches(0.38))))
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": color}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))

    def page_number(slide, n):
        textbox(slide, Inches(12.0), Inches(6.95), Inches(1.0), Inches(0.35), str(n), 11,
                _mix(th["text"], th["bg"], 0.55), align=PP_ALIGN.RIGHT)

    def slide_title(slide, text):
        rect(slide, Inches(0.6), Inches(0.62), Inches(0.1), Inches(0.75), th["accent"])
        textbox(slide, Inches(0.85), Inches(0.5), Inches(11.8), Inches(1.0), _s(text), title_size(_s(text)),
                th["title"], bold=True, anchor=MSO_ANCHOR.MIDDLE)

    for n, s in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        background(slide, th["bg"])

        if s.layout == "title":
            rect(slide, 0, 0, Inches(0.5), Inches(7.5), th["accent"])
            textbox(slide, Inches(1.2), Inches(2.2), Inches(11.2), Inches(1.9), _s(s.title), title_size(_s(s.title), 46),
                    th["title"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
            rect(slide, Inches(1.25), Inches(4.3), Inches(1.7), Inches(0.08), th["accent"])
            if s.subtitle:
                textbox(slide, Inches(1.2), Inches(4.6), Inches(11.2), Inches(1.4), _s(s.subtitle), 22,
                        _mix(th["text"], th["bg"], 0.2))
        elif s.layout == "section":
            background(slide, th["accent"])
            fg = "FFFFFF" if _lum(th["accent"]) < 0.6 else "111111"
            textbox(slide, Inches(1.0), Inches(2.6), Inches(11.3), Inches(2.2), _s(s.title), title_size(_s(s.title), 40),
                    fg, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        elif s.layout == "table" and s.table is not None:
            slide_title(slide, s.title)
            t = s.table
            ncols = max(len(t.headers), 1)
            nrows = len(t.rows) + 1
            fsz = 14 if ncols <= 4 and nrows <= 9 else 12 if ncols <= 6 else 10
            width = Inches(11.9)
            height = min(Inches(0.42) * nrows, Inches(5.2))
            gt = slide.shapes.add_table(nrows, ncols, Inches(0.7), Inches(1.7), width, height).table
            weights = [max(6, min(40, max([len(_s(t.headers[c]))] + [len(_s(r[c])) for r in t.rows if c < len(r)]))) for c in range(ncols)]
            total = sum(weights)
            for c in range(ncols):
                gt.columns[c].width = int(width * weights[c] / total)
            tint = _mix(th["accent"], th["bg"], 0.88)
            on_accent = "FFFFFF" if _lum(th["accent"]) < 0.6 else "111111"
            for ri in range(nrows):
                gt.rows[ri].height = Inches(0.42)
                for ci in range(ncols):
                    cell = gt.cell(ri, ci)
                    raw = t.headers[ci] if ri == 0 else (t.rows[ri - 1][ci] if ci < len(t.rows[ri - 1]) else "")
                    txt = _s(raw)
                    txt = txt if len(txt) <= 140 else txt[:137] + "…"
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(th["accent"] if ri == 0 else (tint if ri % 2 == 0 else th["bg"]))
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    cell.margin_top = cell.margin_bottom = Inches(0.04)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = cell.text_frame.paragraphs[0]
                    runs(p, txt, fsz, on_accent if ri == 0 else th["text"], bold=(ri == 0))
        elif s.layout == "image" and s.image:
            has_title = bool(s.title)
            if has_title:
                slide_title(slide, s.title)
            top = Inches(1.7 if has_title else 0.5)
            max_w, max_h = Inches(11.9), Inches(7.5) - top - Inches(0.6)
            try:
                iw, ih = _img_size(s.image)
                k = min(max_w / iw, max_h / ih)
                w, h = int(iw * k), int(ih * k)
                slide.shapes.add_picture(io.BytesIO(s.image), int((Inches(13.333) - w) / 2), int(top + (max_h - h) / 2), w, h)
            except Exception:
                doc.notes.append("Одно из изображений не удалось вставить в слайд.")
        else:  # bullets
            slide_title(slide, s.title)
            box = slide.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(11.6), Inches(5.0))
            tf = box.text_frame
            tf.word_wrap = True
            chars = sum(len(b) for b in s.bullets) + 25 * len(s.bullets)
            size = 26 if chars < 350 else 22 if chars < 550 else 19 if chars < 800 else 16
            for i, b in enumerate(s.bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(size * 0.55)
                runs(p, b, size, th["text"])
                bullet_para(p, th["accent"])

        if s.layout not in ("title",):
            page_number(slide, n)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = _s(s.notes)
    prs.save(path)


# ============================ PNG ============================


def write_png(doc: Doc, job_dir: Path, stem: str) -> list[Path]:
    from app.convert.render import pdf_to_pngs

    blocks = blocks_of(doc)
    # одна картинка или график -> сам PNG
    if len(blocks) == 1 and blocks[0].kind == "image" and blocks[0].image and not doc.slides:
        out = job_dir / f"{stem}.png"
        Image.open(io.BytesIO(blocks[0].image)).convert("RGBA" if _has_alpha(blocks[0].image) else "RGB").save(out, "PNG")
        return [out]
    tmp_pdf = job_dir / f"{stem}__tmp.pdf"
    write_pdf(doc, tmp_pdf)
    try:
        files, total = pdf_to_pngs(tmp_pdf, job_dir, stem)
    finally:
        tmp_pdf.unlink(missing_ok=True)
    if total > len(files):
        doc.notes.append(f"Показаны первые {len(files)} страниц из {total}.")
    return files


def _has_alpha(data: bytes) -> bool:
    with Image.open(io.BytesIO(data)) as im:
        return im.mode in ("RGBA", "LA")


# ============================ вход ============================


def write(doc: Doc, target: str, job_dir: Path, stem: str) -> list[Path]:
    job_dir.mkdir(parents=True, exist_ok=True)
    if target == "png":
        return write_png(doc, job_dir, stem)
    path = job_dir / f"{stem}.{target}"
    {"xlsx": write_xlsx, "docx": write_docx, "pdf": write_pdf, "pptx": write_pptx}[target](doc, path)
    return [path]
