"""Быстрая проверка без реальных ключей: python -m tests.test_smoke"""
import asyncio
import io
import os
import tempfile
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("ANTHROPIC_API_KEY", "test")
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123:test")

from PIL import Image  # noqa: E402

from app.core import agent, memory  # noqa: E402
from app.media.images import image_block  # noqa: E402
from app.tools import files  # noqa: E402
from app.tools.base import ToolContext  # noqa: E402
from app.tools.web import _is_public_host, fetch_url  # noqa: E402

TABLES = [{"name": "Цены", "headers": ["Товар", "Цена"], "rows": [["Кофе ☕", 120.5], ["=SUM(1,1)", 3]]}]
TEXT = "# Заголовок\n\nАбзац с **жирным** & <символами>.\n\n- пункт один\n- пункт два"
CHART = {"type": "bar", "labels": ["Янв", "Фев", "Мар"], "series": [{"name": "Продажи", "values": [3, 5, 4]}]}


SLIDES = [
    dict(layout="title", title="Обзор", subtitle="Тест"),
    dict(layout="bullets", title="Пункты", bullets=["раз", "два"], notes="заметка"),
    dict(layout="table", title="Таблица", table=TABLES[0]),
    dict(layout="image", title="График", chart=CHART),
]


async def test_files():
    # формат выбирает пользователь и приходит в ctx.target (в make_file параметра format нет)
    cases = [
        ("xlsx", dict(title="Тест", tables=TABLES), ".xlsx"),
        ("pdf", dict(title="Отчёт", content=TEXT, tables=TABLES), ".pdf"),
        ("docx", dict(title="Отчёт", content=TEXT, tables=TABLES), ".docx"),
        ("png", dict(title="График", chart=CHART), ".png"),
        ("png", dict(title="Карточка", content=TEXT), ".png"),
        ("pptx", dict(title="Презентация", slides=SLIDES, design=dict(theme="dark", accent_color="D9480F")), ".pptx"),
    ]
    made = []
    for target, args, ext in cases:
        ctx = ToolContext(user_id=1, target=target)
        out = await files.make_file(ctx, **args)
        assert "Ошибка" not in out and "Недопустимо" not in out, (target, out)
        assert len(ctx.files) == 1 and ctx.files[0].suffix == ext, (target, ctx.files)
        assert ctx.files[0].stat().st_size > 500, ctx.files[0]
        made.append(ctx.files[0])
    # xlsx: строка "=SUM" должна остаться текстом, а не формулой
    from openpyxl import load_workbook
    ws = load_workbook(made[0]).active
    assert ws["A3"].value == "=SUM(1,1)" and ws["A3"].data_type == "s"
    # Excel без таблиц: файл не создаётся, Claude получает указание отказать
    ctx = ToolContext(user_id=1, target="xlsx")
    bad = await files.make_file(ctx, title="x", tables=[])
    assert "Недопустимо" in bad and not ctx.files
    # режим «текст»: файл не создаётся
    ctx = ToolContext(user_id=1, target="text")
    assert "не нужно" in await files.make_file(ctx, title="x", content="y") and not ctx.files
    print("files OK:", [p.name for p in made])


async def test_ssrf():
    assert not _is_public_host("localhost")
    assert not _is_public_host("127.0.0.1")
    assert not _is_public_host("169.254.169.254")
    assert "недоступен" in await fetch_url(ToolContext(1), "http://127.0.0.1:8000/")
    assert "http" in await fetch_url(ToolContext(1), "file:///etc/passwd")
    print("ssrf OK")


def _resp(stop, content):
    return SimpleNamespace(stop_reason=stop, content=content)


async def test_agent():
    memory.DB_PATH = Path(tempfile.mkdtemp()) / "t.db"
    await memory.init_db()

    seen = []
    script = [
        _resp("tool_use", [
            SimpleNamespace(type="text", text="Делаю файл"),
            SimpleNamespace(type="tool_use", id="t1", name="make_file",
                            input=dict(format="xlsx", title="Тест", tables=TABLES)),
        ]),
        _resp("end_turn", [SimpleNamespace(type="text", text="Файл готов.")]),
    ]

    async def fake_create(**kw):
        seen.append(kw["messages"])
        return script[len(seen) - 1]

    agent.client = SimpleNamespace(messages=SimpleNamespace(create=fake_create))

    img = io.BytesIO()
    Image.new("RGB", (3000, 2000), "red").save(img, format="PNG")
    block = image_block(img.getvalue())
    assert block["source"]["media_type"] == "image/jpeg"

    res = await agent.run_agent(42, "сделай таблицу", [block], "xlsx")
    assert res.text == "Файл готов." and len(res.files) == 1 and not res.rejected
    # второй вызов получил tool_result
    assert seen[1][-1]["content"][0]["type"] == "tool_result"

    # память: следующий запрос видит прошлый обмен
    script.append(_resp("end_turn", [SimpleNamespace(type="text", text="Ок")]))
    await agent.run_agent(42, "а теперь?")
    assert len(seen[2]) == 3  # user, assistant, новый user
    assert "[отправлены файлы" in seen[2][1]["content"]

    # отказ: фото в Excel -> понятная ошибка, файла нет, цикл останавливается сразу
    n = len(seen)
    script.append(_resp("tool_use", [
        SimpleNamespace(type="tool_use", id="t2", name="reject_request", input=dict(reason="в Excel нельзя разместить фото")),
    ]))
    res = await agent.run_agent(42, "пришли фото кота", [], "xlsx")
    assert res.rejected and not res.files
    assert res.text.startswith("❌ Недопустимый файл к вашему запросу: в Excel нельзя разместить фото")
    assert len(seen) == n + 1

    await memory.clear_history(42)
    assert await memory.load_history(42) == []
    print("agent+memory+reject OK")


async def main():
    await test_files()
    await test_ssrf()
    await test_agent()


if __name__ == "__main__":
    asyncio.run(main())
