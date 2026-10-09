"""Проверка конвертации всех пар форматов: python -m tests.test_convert"""
import io
import os
import tempfile
from pathlib import Path

os.environ.setdefault("ANTHROPIC_API_KEY", "test")
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123:test")

from PIL import Image  # noqa: E402

from app.convert import NeedsAI, convert  # noqa: E402
from app.convert.writers import make_chart_image, write  # noqa: E402
from app.docmodel import Block, Design, Doc, make_table, markdown_to_blocks  # noqa: E402
from app.formats import EXT_KIND, RULES, TARGETS, FormatError  # noqa: E402

TMP = Path(tempfile.mkdtemp())
TABLE = make_table(["Товар", "Цена", "Остаток"], [["Кофе", 120.5, 10], ["Чай «Эрл Грей»", 80, 25], ["=SUM(1,1)", 3, None]], "Прайс")


def png_bytes(color="steelblue", size=(640, 400)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def sample_doc(with_table=True, with_image=True) -> Doc:
    d = Doc(title="Отчёт о продажах")
    d.blocks = markdown_to_blocks(
        "# Введение\n\nЭто **тестовый** документ с кириллицей и символами & < >.\n\n"
        "- первый пункт\n- второй пункт\n\n## Данные\n\n1. шаг один\n2. шаг два\n"
    )
    if with_table:
        d.blocks.append(TABLE)
    if with_image:
        d.blocks.append(Block("image", image=png_bytes()))
    return d


def make_sources() -> dict[str, Path]:
    """Исходные файлы всех видов (docx есть с таблицей и без)."""
    src = {}
    full = sample_doc()
    for ext in ("docx", "xlsx", "pptx", "pdf"):
        p = TMP / f"sample.{ext}"
        write_to(full, ext, p)
        src[ext] = p
    p = TMP / "notable.docx"
    write_to(sample_doc(with_table=False), "docx", p)
    src["docx_notable"] = p
    (TMP / "sample.csv").write_text("name;price\nкофе;120,5\nчай;80\n", encoding="utf-8")
    src["csv"] = TMP / "sample.csv"
    (TMP / "sample.txt").write_text("# Заметка\n\nПросто текст без таблиц.\n- пункт", encoding="utf-8")
    src["txt"] = TMP / "sample.txt"
    (TMP / "photo.jpg").write_bytes(b"")
    Image.new("RGB", (800, 600), "tomato").save(TMP / "photo.jpg", "JPEG")
    src["image"] = TMP / "photo.jpg"
    Image.new("RGB", (800, 600), "seagreen").save(TMP / "pic.png", "PNG")
    src["image_png"] = TMP / "pic.png"
    return src


def write_to(doc: Doc, ext: str, path: Path) -> None:
    from app.convert import writers
    {"docx": writers.write_docx, "xlsx": writers.write_xlsx, "pptx": writers.write_pptx,
     "pdf": writers.write_pdf}[ext](doc, path)


def verify(path: Path) -> None:
    """Файл реально открывается своей библиотекой и не пустой."""
    ext = path.suffix.lower()
    assert path.exists() and path.stat().st_size > 300, f"пустой файл {path}"
    if ext == ".docx":
        from docx import Document
        assert len(Document(str(path)).element.body) > 0
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        assert load_workbook(path).worksheets
    elif ext == ".pptx":
        from pptx import Presentation
        assert len(Presentation(str(path)).slides) > 0
    elif ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            assert len(pdf.pages) > 0
    elif ext == ".png":
        with Image.open(path) as im:
            assert im.size[0] > 100


def expect(kind_key: str, kind: str, target: str, fname: str) -> str:
    """Что должно получиться по правилам: ok / error / ai."""
    if kind == target or (kind == "image" and target == "png" and fname.endswith(".png")):
        return "error"
    rule = RULES[target].get(kind)
    if rule is None:
        return "error"
    if rule == "ai_table":
        return "ai"
    if rule == "tables":
        return "ok" if kind_key not in ("docx_notable",) else "error"
    return "ok"


def run_matrix(label: str) -> None:
    src = make_sources()
    ok = err = ai = 0
    for key, path in src.items():
        kind = EXT_KIND[path.suffix.lower()]
        for target in TARGETS:
            want = expect(key, kind, target, path.name)
            try:
                res = convert(path, kind, target)
                got = "ok"
                for f in res.files:
                    verify(f)
                assert res.files, "нет файлов"
            except NeedsAI:
                got = "ai"
            except FormatError as e:
                got = "error"
                assert e.user_message().startswith("❌ Недопустимый файл к вашему запросу:")
            assert got == want, f"[{label}] {key} -> {target}: ожидали {want}, получили {got}"
            ok += got == "ok"
            err += got == "error"
            ai += got == "ai"
    print(f"matrix [{label}] OK: {ok} конвертаций, {err} отказов, {ai} через Claude")


def test_content_preserved():
    """Данные не теряются при переносе: docx -> xlsx, xlsx -> docx, pptx -> docx, pdf -> docx."""
    src = make_sources()
    from app.convert.readers import read_docx, read_pdf, read_pptx, read_xlsx

    x = convert(src["docx"], "docx", "xlsx").files[0]
    t = read_xlsx(x).tables[0]
    assert t.headers == ["Товар", "Цена", "Остаток"] and t.rows[0][0] == "Кофе", t
    d = convert(src["xlsx"], "xlsx", "docx").files[0]
    assert read_docx(d).tables[0].rows[1][0] == "Чай «Эрл Грей»"
    d2 = convert(src["pptx"], "pptx", "docx").files[0]
    texts = " ".join(b.text for b in read_docx(d2).blocks)
    assert "Введение" in texts and "первый пункт" in texts, texts
    d3 = convert(src["pdf"], "pdf", "docx").files[0]
    texts = " ".join(b.text for b in read_docx(d3).blocks)
    assert "тестовый" in texts, texts
    assert read_pdf(src["pdf"]).tables, "таблица в PDF не распознана"
    assert read_pptx(src["pptx"]).tables
    print("content preserved OK")


def test_pptx_design():
    from pptx import Presentation
    from app.convert.writers import resolve_design

    doc = sample_doc()
    for name, design in {
        "dark": Design(theme="dark"),
        "custom": Design(theme="light", accent="D9480F", background="102030", font="Georgia"),
        "bad": Design(theme="nonexistent", accent="zzz"),
    }.items():
        doc.design = design
        p = TMP / f"design_{name}.pptx"
        write("pptx_dummy" and doc, "pptx", TMP, f"design_{name}")
        assert p.exists()
        assert len(Presentation(str(p)).slides) >= 4
    th = resolve_design(Design(theme="light", background="102030"))
    assert th["bg"] == "102030" and th["text"] == "F3F4F6"          # тёмный фон -> светлый текст
    assert th["title"] == "FFFFFF"
    assert resolve_design(Design(theme="nonexistent", accent="zzz"))["accent"] == "2F5597"
    print("pptx design OK")


def test_chart_png():
    data = make_chart_image({"type": "line", "labels": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}, "T")
    assert data[:4] == b"\x89PNG"
    try:
        make_chart_image({"type": "bar", "labels": ["a", "b"], "series": [{"values": [1]}]})
    except ValueError:
        pass
    else:
        raise AssertionError("несовпадение длин не поймано")
    files = write(Doc(blocks=[Block("image", image=data)]), "png", TMP, "chart_only")
    assert len(files) == 1 and files[0].suffix == ".png"
    print("chart OK")


def main():
    run_matrix("с LibreOffice")
    os.environ["DISABLE_LIBREOFFICE"] = "1"
    run_matrix("без LibreOffice")
    os.environ.pop("DISABLE_LIBREOFFICE")
    test_content_preserved()
    test_pptx_design()
    test_chart_png()


if __name__ == "__main__":
    main()
