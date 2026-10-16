"""Языки и напоминания. Запуск: python -m tests.test_i18n"""
import asyncio
import html
import os
import re
import tempfile
from datetime import timedelta

os.environ["GA4_API_SECRET"] = ""  # тесты не шлют события в настоящий GA4
os.environ["DATA_DIR"] = tempfile.mkdtemp()

from app import config, db, i18n, reminders  # noqa: E402
from app.bot import welcome  # noqa: E402
from app.niches import NICHES  # noqa: E402

FIELDS = lambda s: sorted(re.findall(r"\{(\w+)\}", s))  # noqa: E731


def test_texts_complete():
    for key, v in i18n.TEXTS.items():
        assert set(v) == set(i18n.LANGS), key
        for lg in i18n.LANGS:                       # одинаковые плейсхолдеры и нет пустых строк
            assert v[lg].strip(), (key, lg)
            assert FIELDS(v[lg]) == FIELDS(v["ru"]), (key, lg)
    for table in (i18n.MONTHS, i18n.WEEKDAYS, i18n.WEEKDAYS2, i18n.COMMANDS, i18n.COMMANDS_DEMO):
        assert set(table) == set(i18n.LANGS)
    assert all(len(m) == 13 for m in i18n.MONTHS.values())
    assert all(len(w) == 7 for w in i18n.WEEKDAYS.values()) and all(len(w) == 7 for w in i18n.WEEKDAYS2.values())
    for lg in i18n.LANGS:
        for c, d in i18n.COMMANDS[lg] + i18n.COMMANDS_DEMO[lg]:
            assert re.fullmatch(r"[a-z_]{1,32}", c) and 3 <= len(d) <= 256


def test_niches_translated():
    for k, n in NICHES.items():
        for lg in ("uk", "en"):
            v = i18n.niche_view(k, lg)
            assert len(v.services) == len(n.services) and len(v.samples) == len(n.samples), (k, lg)
            assert v.company and v.about and v.address and v.hours and v.photo_example
        for lg in i18n.LANGS:
            for demo in (True, False):
                config.DEMO_MODE = demo
                w = welcome(k, lg)
                assert len(w) < 4096, (k, lg)
                assert all(html.escape(name) in w for name in i18n.niche_view(k, lg).services)
    config.DEMO_MODE = True
    assert i18n.service_label("salon", "Женская стрижка", "en") == "Women's haircut"
    assert i18n.service_label("salon", "нет такой", "en") == "нет такой"
    assert i18n.find_service_any("salon", "Women's haircut")[0] == 0
    assert i18n.find_service_any("salon", "жіноча стрижка")[0] == 0
    assert i18n.find_service_any("salon", "Женская стрижка")[0] == 0
    assert i18n.find_service_any("salon", "rocket") is None and i18n.find_service_any("salon", "") is None


def test_resolve():
    assert i18n.default_lang() == "uk"                                       # по умолчанию украинский
    assert i18n.resolve({"lang": "en", "tg_lang": "uk"}) == "en"            # сохранённый язык главнее всего
    assert i18n.resolve({"tg_lang": "ru"}) == "uk" and i18n.resolve(None, "en") == "uk"   # Telegram не учитывается
    assert i18n.resolve({"lang": "xx"}) == "uk" and i18n.resolve(None) == "uk"
    assert not hasattr(i18n, "guess")


def test_detect():
    d = i18n.detect
    assert d("Привіт, хочу записатись на стрижку") == "uk" and d("Скільки коштує манікюр?") == "uk"
    assert d("Добрий день") == "uk" and d("дякую!") == "uk"
    assert d("Здравствуйте, хочу записаться на завтра") == "ru" and d("Сколько стоит маникюр?") == "ru"
    assert d("привет") == "ru" and d("Что по ценам?") == "ru"
    assert d("I want to book a haircut for tomorrow") == "en" and d("hello") == "en" and d("How much is it?") == "en"
    # нейтральное: язык не меняется
    for neutral in ("", None, "   ", "050 123 45 67", "5", "👍", "Anna", "Anna Ivanova", "/cabinet", "/start x", "ок"):
        assert d(neutral) is None, neutral


def test_funnel_texts_en():
    import re
    from app.funnel_texts import FT, ft
    for key, v in FT.items():
        assert set(v) == {"uk", "ru", "en"}, key
        for lg in ("ru", "en"):
            a, b = v["uk"], v[lg]
            if isinstance(a, str):
                assert sorted(re.findall(r"\{\w+\}", a)) == sorted(re.findall(r"\{\w+\}", b)), (key, lg)
            elif isinstance(a, list):
                assert len(a) == len(b), (key, lg)
            else:
                assert set(a) == set(b), (key, lg)
    assert ft("en", "btn_book").startswith("📅") and ft("xx", "btn_book") == ft("uk", "btn_book")


def test_report_langs():
    import io
    from openpyxl import load_workbook
    from app.report import leads_to_xlsx
    lead = dict(id=1, created="2026-10-15T10:00:00", niche="salon", service="Женская стрижка", price=700,
                date="2026-10-16", time="10:00", name="Anna", phone="050", comment=None, source="AI", status="ok")
    for lg, head, status, svc in (("uk", "Послуга", "Підтверджена", "Жіноча стрижка"),
                                  ("ru", "Услуга", "Подтверждена", "Женская стрижка"),
                                  ("en", "Service", "Confirmed", "Women's haircut")):
        ws = load_workbook(io.BytesIO(leads_to_xlsx([lead], lg))).active
        assert ws["D1"].value == head and ws["L2"].value == status and ws["D2"].value == svc, lg
    ws = load_workbook(io.BytesIO(leads_to_xlsx([lead]))).active          # по умолчанию украинский
    assert ws["D1"].value == "Послуга"


def test_no_language_selection():
    for lg in i18n.LANGS:
        assert "/language" not in i18n.t(lg, "profile") and "/language" not in i18n.t(lg, "cab_todo")
        assert "/language" not in i18n.t(lg, "welcome", company="x", about="x", services="x", hours="x", address="x",
                                          photo="x", examples="x")
    for key in ("lang_choose", "lang_set_first", "lang_changed", "lang_tool_ok"):
        assert key not in i18n.TEXTS
    for cmds in i18n.COMMANDS.values():
        assert "language" not in [c for c, _ in cmds]


class FakeBot:
    def __init__(self, fail=False):
        self.sent, self.fail = [], fail

    async def send_message(self, chat_id, text):
        if self.fail:
            raise RuntimeError("blocked")
        self.sent.append((chat_id, text))


async def test_reminders():
    await db.init_db()
    now = db.now_tz().replace(microsecond=0)

    def at(delta):
        d = now + delta
        return d.date().isoformat(), d.strftime("%H:%M")

    await db.touch_user(1, "a")
    await db.touch_user(2, "b")
    await db.set_lang(2, "en")

    async def add(uid, delta, status=None):
        d, t = at(delta)
        lid = await db.add_lead(uid, "x", "salon", "Женская стрижка", 700, d, t, "Анна", "0501234567")
        if status:
            await db.set_status(lid, status)
        return lid

    far = await add(1, timedelta(days=3))                   # рано: напоминаний ещё нет
    day = await add(2, timedelta(hours=23, minutes=30))     # создана впритык (<25ч): сутки пропускаем, 2ч придёт
    soon = await add(1, timedelta(hours=1, minutes=30))     # <3ч: ждёт только... ничего, обе отметки уже стоят
    cancelled = await add(1, timedelta(hours=10), "cancel")
    declined = await add(1, timedelta(hours=10), "no")
    for lid, r24, r2 in ((far, 0, 0), (day, 1, 0), (soon, 1, 1)):
        l = await db.get_lead(lid)
        assert (l["rem24"], l["rem2"]) == (r24, r2), (lid, l["rem24"], l["rem2"])

    bot = FakeBot()
    assert await reminders.run_once(bot, now) == 0          # сейчас слать нечего
    # прошло ~21.5 ч: для day осталось 2 часа - придёт напоминание за 2 часа на английском
    assert await reminders.run_once(bot, now + timedelta(hours=21, minutes=45)) == 1
    chat, text = bot.sent[0]
    assert chat == 2 and "in a couple of hours" in text and html.escape("Women's haircut") in text
    assert await reminders.run_once(bot, now + timedelta(hours=21, minutes=50)) == 0     # без дублей
    # для far за 24ч: украинский текст, потом за 2ч, и 2h вытесняет 24h, если цикл проспал оба срока
    t24 = now + timedelta(days=2, hours=1)
    assert await reminders.run_once(bot, t24) == 1
    assert bot.sent[-1][0] == 1 and "Нагадуємо" in bot.sent[-1][1] and f"№{far}" in bot.sent[-1][1]
    assert await reminders.run_once(bot, t24 + timedelta(minutes=5)) == 0
    assert await reminders.run_once(bot, now + timedelta(days=3) - timedelta(hours=1)) == 1
    assert "Через пару годин" in bot.sent[-1][1]
    # цикл спал: сразу осталось 1.5 часа, приходит только одно напоминание (за 2 часа)
    l = await add(1, timedelta(days=5))
    n = len(bot.sent)
    assert await reminders.run_once(bot, now + timedelta(days=5) - timedelta(hours=1, minutes=30)) == 1
    assert len(bot.sent) == n + 1 and "Через пару годин" in bot.sent[-1][1]
    # отменённые и отклонённые не напоминаются; ошибка отправки не повторяется
    for lid in (cancelled, declined):
        assert all(f"№{lid}" not in t for _, t in bot.sent)
    bad = FakeBot(fail=True)
    z = await add(2, timedelta(days=7))
    assert await reminders.run_once(bad, now + timedelta(days=7) - timedelta(hours=1)) == 0
    good = FakeBot()
    assert await reminders.run_once(good, now + timedelta(days=7) - timedelta(minutes=50)) == 0 and not good.sent
    assert (await db.get_lead(z))["rem2"] == 1


def main():
    test_texts_complete()
    test_niches_translated()
    test_resolve()
    test_detect()
    test_funnel_texts_en()
    test_report_langs()
    test_no_language_selection()
    asyncio.run(test_reminders())
    print("языки и напоминания: все проверки пройдены")


if __name__ == "__main__":
    main()
