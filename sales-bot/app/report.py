"""Excel-отчёт по заявкам: то, что получает владелец (на русском) или клиент в демо (на его языке)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import i18n
from .niches import NICHES

STATUS = {
    "ru": {"new": "Новая", "ok": "Подтверждена", "no": "Отклонена", "cancel": "Отменена клиентом"},
    "uk": {"new": "Нова", "ok": "Підтверджена", "no": "Відхилена", "cancel": "Скасована клієнтом"},
    "en": {"new": "New", "ok": "Confirmed", "no": "Declined", "cancel": "Cancelled by client"},
}
HEAD = {
    "ru": ["№", "Создана", "Компания", "Услуга", "Цена, ₴", "Дата", "Время", "Клиент", "Телефон",
           "Комментарий", "Источник", "Статус", "Запись в CRM"],
    "uk": ["№", "Створена", "Компанія", "Послуга", "Ціна, ₴", "Дата", "Час", "Клієнт", "Телефон",
           "Коментар", "Джерело", "Статус", "Запис у CRM"],
    "en": ["#", "Created", "Company", "Service", "Price, ₴", "Date", "Time", "Client", "Phone",
           "Comment", "Source", "Status", "CRM record"],
}
SHEET = {"ru": "Заявки", "uk": "Заявки", "en": "Bookings"}
TOTAL = {"ru": "Сумма, ₴:", "uk": "Сума, ₴:", "en": "Total, ₴:"}


def leads_to_xlsx(leads: list, lang: str = "uk") -> bytes:
    lang = lang if lang in i18n.LANGS else i18n.default_lang()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET[lang]
    ws.append(HEAD[lang])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center")
    for l in leads:
        k = l["niche"]
        company = i18n.niche_view(k, lang).company if k in NICHES else k
        service = i18n.service_label(k, l["service"], lang) if k in NICHES else l["service"]
        ws.append([l["id"], l["created"].replace("T", " "), company, service, l["price"],
                   l["date"], l["time"], l["name"], l["phone"], l["comment"] or "",
                   l["source"], STATUS[lang].get(l["status"], l["status"]),
                   f"{l['crm']} №{l['external_id']}" if l.get("external_id") else ""])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A2"
    if leads:
        ws.append([])
        ws.append(["", "", "", TOTAL[lang], f"=SUM(E2:E{len(leads) + 1})"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
